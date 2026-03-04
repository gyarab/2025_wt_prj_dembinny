"""
accounts/views_admin.py
───────────────────────
System-Admin-only management views for Classes, FundGroups, and ClassMemberships.

All views are protected by @admin_required.

Views
─────
admin_panel_view          – overview: all classes + their groups and members
fund_group_create_view    – create a new FundGroup for a class
fund_group_edit_view      – edit an existing FundGroup's name / permissions
fund_group_delete_view    – remove a FundGroup (POST-only; blocked if members assigned)
membership_create_view    – assign a treasurer to a class in a FundGroup
membership_edit_view      – change which FundGroup a membership belongs to
membership_delete_view    – remove a ClassMembership (POST-only)
"""

import io

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from django.contrib import messages
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render

from .decorators import admin_required
from .forms_admin import ClassMembershipForm, FundGroupForm
from .models import ClassMembership, FundGroup, SchoolClass


# ── Overview ──────────────────────────────────────────────────────────────────

@admin_required
def admin_panel_view(req):
    """
    Main admin panel page.
    Lists every SchoolClass with its FundGroups and ClassMembership entries.
    """
    classes = (
        SchoolClass.objects
        .prefetch_related('fund_groups__memberships__user', 'memberships__user',
                          'memberships__fund_group')
        .select_related('teacher')
        .order_by('name')
    )
    return render(req, 'accounts/admin_panel.html', {'classes': classes})


# ── FundGroup CRUD ────────────────────────────────────────────────────────────

@admin_required
def fund_group_create_view(req):
    """
    GET  → show blank FundGroup form (optionally pre-filled from ?class=<pk>).
    POST → validate and create the FundGroup.
    """
    school_class = None
    if req.GET.get('class') or req.POST.get('school_class'):
        pk = req.POST.get('school_class') or req.GET.get('class')
        try:
            school_class = SchoolClass.objects.get(pk=int(pk))
        except (SchoolClass.DoesNotExist, ValueError):
            pass

    if req.method == 'POST':
        form = FundGroupForm(req.POST, school_class=school_class)
        if form.is_valid():
            group = form.save()
            messages.success(req, f'✅ Group "{group.name}" created for {group.school_class}.')
            return redirect('admin_panel')
        messages.error(req, 'Please fix the errors below.')
    else:
        form = FundGroupForm(school_class=school_class)

    return render(req, 'accounts/fund_group_form.html', {
        'form':         form,
        'school_class': school_class,
        'title':        'Create Permission Group',
        'submit_label': 'Create Group',
    })


@admin_required
def fund_group_edit_view(req, group_id):
    """
    GET  → pre-filled form for an existing FundGroup.
    POST → update name and permission flags.
    """
    group = get_object_or_404(FundGroup, pk=group_id)

    if req.method == 'POST':
        form = FundGroupForm(req.POST, instance=group, school_class=group.school_class)
        if form.is_valid():
            form.save()
            messages.success(req, f'✅ Group "{group.name}" updated.')
            return redirect('admin_panel')
        messages.error(req, 'Please fix the errors below.')
    else:
        form = FundGroupForm(instance=group, school_class=group.school_class)

    return render(req, 'accounts/fund_group_form.html', {
        'form':         form,
        'group':        group,
        'school_class': group.school_class,
        'title':        f'Edit Group — {group.name}',
        'submit_label': 'Save Changes',
    })


@admin_required
def fund_group_delete_view(req, group_id):
    """POST-only: remove a FundGroup if no memberships are using it."""
    if req.method != 'POST':
        return redirect('admin_panel')

    group = get_object_or_404(FundGroup, pk=group_id)

    if group.memberships.exists():
        count = group.memberships.count()
        messages.error(
            req,
            f'Cannot delete "{group.name}" — {count} member(s) still assigned to it. '
            f'Move or remove them first.'
        )
        return redirect('admin_panel')

    name  = group.name
    klass = str(group.school_class)
    group.delete()
    messages.success(req, f'🗑 Deleted group "{name}" from {klass}.')
    return redirect('admin_panel')


# ── ClassMembership CRUD ──────────────────────────────────────────────────────

@admin_required
def membership_create_view(req):
    """
    GET  → show blank membership form (optionally pre-fill class from ?class=<pk>).
    POST → validate and create the ClassMembership.
    """
    school_class = None
    if req.GET.get('class') or req.POST.get('school_class'):
        pk = req.POST.get('school_class') or req.GET.get('class')
        try:
            school_class = SchoolClass.objects.get(pk=int(pk))
        except (SchoolClass.DoesNotExist, ValueError):
            pass

    if req.method == 'POST':
        form = ClassMembershipForm(req.POST, school_class=school_class)
        if form.is_valid():
            membership = form.save()
            messages.success(
                req,
                f'✅ {membership.user.get_full_name() or membership.user.username} '
                f'added to {membership.school_class}'
                + (f' as {membership.fund_group.name}.' if membership.fund_group else '.')
            )
            return redirect('admin_panel')
        messages.error(req, 'Please fix the errors below.')
    else:
        form = ClassMembershipForm(school_class=school_class)

    return render(req, 'accounts/membership_form.html', {
        'form':         form,
        'school_class': school_class,
        'title':        'Add Class Member',
        'submit_label': 'Add Member',
    })


@admin_required
def membership_edit_view(req, membership_id):
    """
    GET  → pre-filled form for an existing ClassMembership.
    POST → update the fund_group (class and user are read-only after creation).
    """
    membership = get_object_or_404(ClassMembership, pk=membership_id)

    if req.method == 'POST':
        form = ClassMembershipForm(req.POST, instance=membership,
                                   school_class=membership.school_class)
        if form.is_valid():
            form.save()
            messages.success(
                req,
                f'✅ Updated: {membership.user.get_full_name() or membership.user.username} '
                f'in {membership.school_class}'
                + (f' → {membership.fund_group.name}.' if membership.fund_group else '.')
            )
            return redirect('admin_panel')
        messages.error(req, 'Please fix the errors below.')
    else:
        form = ClassMembershipForm(instance=membership, school_class=membership.school_class)

    return render(req, 'accounts/membership_form.html', {
        'form':         form,
        'membership':   membership,
        'school_class': membership.school_class,
        'title':        'Edit Class Membership',
        'submit_label': 'Save Changes',
    })


@admin_required
def membership_delete_view(req, membership_id):
    """POST-only: remove a ClassMembership."""
    if req.method != 'POST':
        return redirect('admin_panel')

    membership = get_object_or_404(ClassMembership, pk=membership_id)
    name  = membership.user.get_full_name() or membership.user.username
    klass = str(membership.school_class)
    membership.delete()
    messages.success(req, f'🗑 Removed {name} from {klass}.')
    return redirect('admin_panel')


# ── Credentials Export ────────────────────────────────────────────────────────

@admin_required
def export_credentials_view(req, class_id):
    """
    Generate and download an Excel (.xlsx) sheet containing the pre-generated
    usernames and passwords for all students in *class_id*.

    Only rows where a plain-text password was recorded at import time are
    included (i.e. newly created accounts from the CSV importer).
    """
    from importer.models import ImportRow

    school_class = get_object_or_404(SchoolClass, pk=class_id)

    rows = (
        ImportRow.objects
        .filter(batch__school_class=school_class)
        .exclude(plain_password='')
        .values('last_name', 'first_name', 'username', 'plain_password', 'variable_symbol')
        .order_by('last_name', 'first_name')
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Credentials'

    # Header row styling
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(fill_type='solid', fgColor='2D6A4F')
    header_alignment = Alignment(horizontal='center', vertical='center')

    headers = ['Last Name', 'First Name', 'Username', 'Password', 'Variable Symbol']
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    # Data rows
    for row_idx, row in enumerate(rows, start=2):
        ws.cell(row=row_idx, column=1, value=row['last_name'])
        ws.cell(row=row_idx, column=2, value=row['first_name'])
        ws.cell(row=row_idx, column=3, value=row['username'])
        ws.cell(row=row_idx, column=4, value=row['plain_password'])
        ws.cell(row=row_idx, column=5, value=row['variable_symbol'])

    # Auto-fit column widths
    for col in ws.columns:
        max_len = max((len(str(cell.value or '')) for cell in col), default=0)
        ws.column_dimensions[col[0].column_letter].width = max(max_len + 4, 14)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    safe_name = ''.join(c if c.isalnum() else '_' for c in school_class.name)
    filename = f'credentials_{safe_name}.xlsx'

    response = HttpResponse(
        buf.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response
